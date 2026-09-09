from pathlib import Path
import json
import re
import unicodedata
import urllib.request
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / ".historical-tmp"
OUT = ROOT / "historical-data.js"
TMP.mkdir(exist_ok=True)

SOURCES = {
    "2023-2024": "https://raw.githubusercontent.com/LBrontesi/Applicazione-fantacalcio/refs/heads/main/Statistiche_Fantacalcio_Stagione_2023_24.xlsx",
    "2024-2025": "https://raw.githubusercontent.com/LBrontesi/Applicazione-fantacalcio/refs/heads/main/Statistiche_Fantacalcio_Stagione_2024_25.xlsx",
    "2025-2026": "https://raw.githubusercontent.com/e-manganelli/Fantabot/refs/heads/main/Copy%20of%20Statistiche_Fantacalcio_Stagione_2025_26.xlsx",
}

ALIASES = {
    "id": {"id", "codice", "codicegiocatore", "playerid"},
    "role": {"r", "ruolo", "role"},
    "name": {"nome", "giocatore", "name"},
    "team": {"squadra", "team", "club"},
    "pv": {"pv", "presenze", "presenzaconvoto", "pres"},
    "mv": {"mv", "mediavoto", "mediavoto"},
    "fm": {"fm", "fantamedia", "fantavotomedio"},
    "gf": {"gf", "gol", "golfatti", "g"},
    "gs": {"gs", "golsubiti"},
    "rp": {"rp", "rigoriparati"},
    "rc": {"rc", "rigoricalciati", "rigoritirati"},
    "rg": {"r+", "rigorisegnati", "rigorirealizzati"},
    "rs": {"r-", "rigorisbagliati"},
    "ass": {"ass", "assist", "a"},
    "amm": {"amm", "ammonizioni"},
    "esp": {"esp", "espulsioni"},
    "au": {"au", "autogol"},
}


def norm(value):
    s = str(value or "").strip().lower()
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9+%-]+", "", s)


def number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, str):
        value = value.replace(",", ".").replace("%", "")
    try:
        x = float(value)
        return int(x) if x.is_integer() else x
    except Exception:
        return None


def fetch(url, target):
    if not target.exists() or target.stat().st_size < 1000:
        req = urllib.request.Request(url, headers={"User-Agent": "AstaLab historical-data builder"})
        with urllib.request.urlopen(req, timeout=60) as r, open(target, "wb") as f:
            f.write(r.read())


def find_col(headers, aliases):
    for i, h in enumerate(headers):
        if h in aliases:
            return i
    return None


def parse_xlsx(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    headers = [norm(x) for x in header]
    cols = {k: find_col(headers, a) for k, a in ALIASES.items()}

    required = ["role", "name", "pv", "gf", "ass"]
    missing = [k for k in required if cols[k] is None]
    if missing:
        raise RuntimeError(f"{path.name}: colonne mancanti: {missing}; trovate: {header}")

    out = []
    for row in rows:
        if not any(v is not None and str(v).strip() for v in row):
            continue
        def get(k):
            i = cols[k]
            return row[i] if i is not None and i < len(row) else None
        role = str(get("role") or "").strip().upper()
        name = str(get("name") or "").strip()
        if role not in {"P", "D", "C", "A"} or not name:
            continue
        out.append({
            "id": number(get("id")),
            "role": role,
            "name": name,
            "team": str(get("team") or "").strip(),
            "pv": number(get("pv")) or 0,
            "mv": number(get("mv")),
            "fm": number(get("fm")),
            "gf": number(get("gf")) or 0,
            "gs": number(get("gs")),
            "rp": number(get("rp")),
            "rc": number(get("rc")),
            "rg": number(get("rg")),
            "rs": number(get("rs")),
            "ass": number(get("ass")) or 0,
            "amm": number(get("amm")),
            "esp": number(get("esp")),
            "au": number(get("au")),
        })
    return out

seasons = {}
for season, url in SOURCES.items():
    path = TMP / f"{season}.xlsx"
    fetch(url, path)
    seasons[season] = parse_xlsx(path)

payload = {
    "version": "1.0.0",
    "generatedAt": __import__("datetime").datetime.utcnow().isoformat(timespec="seconds") + "Z",
    "source": "Fantacalcio statistics datasets; source files mirrored in public GitHub repositories",
    "seasonWeights": {"2023-2024": 0.20, "2024-2025": 0.30, "2025-2026": 0.50},
    "metrics": {
        "goals": "gf / pv",
        "assists": "ass / pv",
        "mv": "stored for analysis; not part of the current goal/assist score",
        "fm": "stored for analysis; not part of the current goal/assist score",
    },
    "seasons": seasons,
}

OUT.write_text(
    "const HISTORICAL_DATA=" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\nwindow.HISTORICAL_DATA=HISTORICAL_DATA;\n",
    encoding="utf-8",
)

print("Historical dataset generated:")
for season, rows in seasons.items():
    print(f"  {season}: {len(rows)} players")
